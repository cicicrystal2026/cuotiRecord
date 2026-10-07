import io
import time
from datetime import datetime, timedelta, timezone
import fitz
import pytest
from PIL import Image
from docx import Document
from docx.oxml import parse_xml
from fastapi.testclient import TestClient
from app import storage as db, config, domain, ai
from app.main import app

HEADERS={'X-Cuoti-Client':'local-v1'}

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(db,'DATA_DIR',tmp_path)
    monkeypatch.setattr(db,'DB_PATH',tmp_path/'library.db')
    monkeypatch.setattr(config,'DATA_DIR',tmp_path)
    monkeypatch.setattr(config,'CONFIG_FILE',tmp_path/'settings.json')
    monkeypatch.delenv('DASHSCOPE_API_KEY',raising=False)
    with TestClient(app,headers=HEADERS) as c:
        c.put('/api/profile',json={'nickname':'同学','grade':'高一','curriculum':'人教A版','semester':'2026秋','start_date':'2026-09-01','end_date':'2027-01-31'})
        yield c


def upload(c, filename='test.jpg', content=None):
    if content is None:
        b=io.BytesIO();Image.new('RGB',(900,1200),'white').save(b,'JPEG');content=b.getvalue()
    r=c.post('/api/sources',data={'name':'测试资料','learning_date':'2026-09-20'},files={'files':(filename,content)})
    assert r.status_code==200,r.text
    sid=r.json()['id']
    for _ in range(100):
        value=c.get('/api/sources/'+sid).json()
        if value['tasks'][0]['status'] not in ('queued','running'):
            assert value['tasks'][0]['status']=='success',value
            return c.get("/api/sources/"+sid).json()
        time.sleep(.1)
    pytest.fail('conversion timed out')


def confirmed(c, source=None, missing=False):
    source=source or upload(c)
    q=c.post('/api/sources/'+source['id']+'/questions',json={'label':'1','stem':'She enjoys ___ (read) books.','work':'' if missing else 'read','missing':missing}).json()
    a=c.post('/api/questions/'+q['id']+'/manual-analysis').json()['analysis']
    body={'analysis_id':a['id'],'question_version':q['version'],'status':'confirmed','answer':'reading；enjoy后接动名词','correct_parts':'','issues':'enjoy后不能用动词原形','evidence':'read','check_action':'检查动词搭配','knowledge_ids':[domain.NODES[0]['id']] if hasattr(domain,'NODES') else ['math.function.monotonicity'],'causes':['语法'],'uncertainty':''}
    r=c.put('/api/questions/'+q['id']+'/analysis',json=body)
    if missing:
        assert r.status_code==400
    else:
        assert r.status_code==200,r.text
        c.post('/api/sources/'+source['id']+'/save')
    return q,body,source


def test_upload_images_pdf_docx(client):
    s=upload(client);assert len(s['pages'])==1
    p=fitz.open();page=p.new_page();page.insert_text((50,50),'Math x^2');data=p.tobytes();p.close()
    assert len(upload(client,'text.pdf',data)['pages'])==1
    b=io.BytesIO();Image.new('RGB',(500,700),'white').save(b,'JPEG')
    p=fitz.open();p.new_page().insert_image(fitz.Rect(0,0,500,700),stream=b.getvalue());data=p.tobytes();p.close()
    assert len(upload(client,'scan.pdf',data)['pages'])==1
    d=Document();d.add_paragraph('高中数学：函数与单调性');d.add_paragraph('f(x)=x²');d.add_table(rows=2,cols=2).cell(0,0).text='题号'
    d.paragraphs[-1]._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'))
    b=io.BytesIO();d.save(b)
    s=upload(client,'formula.docx',b.getvalue());assert len(s['pages'])>=1;assert '重排' in s['warning']
    assert client.get('/api/pages/'+s['pages'][0]['id']+'/image').headers['content-type']=='image/jpeg'


def test_missing_key_and_security(client):
    s=upload(client)
    assert client.post('/api/sources/'+s['id']+'/recognize').status_code==400
    assert client.put('/api/config',json={'base_url':'https://evil.example/v1'}).status_code==400
    assert client.post('/api/sample',headers={'origin':'https://evil.example'}).status_code==403
    assert client.get('/api/health',headers={'host':'evil.example'}).status_code==403
    assert client.post('/api/sample',headers={'X-Cuoti-Client':''}).status_code==403


def test_key_encrypted_no_echo(client):
    r=client.put('/api/config',json={'api_key':'sk-test-0123456789'})
    assert r.status_code==200,r.text
    assert 'sk-test' not in r.text
    assert 'sk-test' not in config.CONFIG_FILE.read_text()
    assert config.settings(True)['api_key']=='sk-test-0123456789'
    client.put('/api/config',json={'clear_key':True})
    assert not client.get('/api/config').json()['configured']


def test_missing_work_cannot_confirm(client):
    q,body,s=confirmed(client,missing=True)
    body['status']='pending'
    r=client.put('/api/questions/'+q['id']+'/analysis',json=body)
    assert r.status_code==200
    a=r.json()['analysis'];assert a['issues']=='' and a['causes']==[]
    client.post('/api/sources/'+s['id']+'/save')
    assert client.get('/api/review/candidates').json()==[]


def test_version_and_evidence_guards(client):
    q,body,s=confirmed(client)
    body['evidence']='不存在的作答'
    assert client.put('/api/questions/'+q['id']+'/analysis',json=body).status_code==400
    r=client.put('/api/questions/'+q['id'],json={'version':1,'stem':'求x²的单调性'})
    assert r.json()['version']==2
    assert client.get('/api/review/candidates').json()==[]
    assert client.put('/api/questions/'+q['id'],json={'version':1,'stem':'旧版'}).status_code==409
    assert client.put('/api/questions/'+q['id']+'/analysis',json=body).status_code==400


def session(c,q,key):
    r=c.post('/api/review/sessions',json={'question_ids':[q['id']],'name':'测试复习','request_key':key})
    assert r.status_code==200,r.text
    return r.json()


def test_answer_hidden_hint_idempotency_pause(client):
    q,body,s=confirmed(client)
    sess=session(client,q,'same-request');item=sess['items'][0]
    assert 'answer' not in item['snapshot']
    assert session(client,q,'same-request')['id']==sess['id']
    assert client.post('/api/review/items/'+item['id']+'/attempt',json={'result':'right'}).status_code==400
    client.post('/api/review/sessions/'+sess['id']+'/state',json={'action':'pause'})
    assert client.get('/api/review/sessions/'+sess['id']).json()['status']=='paused'
    client.post('/api/review/sessions/'+sess['id']+'/state',json={'action':'resume'})
    client.post('/api/review/items/'+item['id']+'/reveal',json={'kind':'hint'})
    r=client.post('/api/review/items/'+item['id']+'/reveal',json={'kind':'answer'})
    assert r.json()['items'][0]['snapshot']['answer']==body['answer']
    assert client.post('/api/review/items/'+item['id']+'/attempt',json={'result':'right'}).status_code==400
    assert client.post('/api/review/items/'+item['id']+'/attempt',json={'result':'hint'}).json()['status']=='completed'
    client.post('/api/review/items/'+item['id']+'/attempt',json={'result':'wrong'})
    assert len(db.rows('SELECT * FROM attempts'))==1
    assert client.get('/api/questions/'+q['id']).json()['state']=='需要巩固'


def test_independent_24h_and_reset(client):
    q,body,s=confirmed(client)
    for n in range(2):
        sess=session(client,q,str(n));item=sess['items'][0]
        client.post('/api/review/items/'+item['id']+'/reveal',json={'kind':'answer'})
        client.post('/api/review/items/'+item['id']+'/attempt',json={'result':'right'})
    assert client.get('/api/questions/'+q['id']).json()['state']=='本次独立做对'
    old=(datetime.now(timezone.utc)-timedelta(hours=25)).isoformat()
    db.execute('UPDATE attempts SET completed_at=? WHERE id=(SELECT id FROM attempts ORDER BY completed_at LIMIT 1)',(old,))
    assert client.get('/api/questions/'+q['id']).json()['state']=='已验证'
    sess=session(client,q,'reset');item=sess['items'][0]
    client.post('/api/review/items/'+item['id']+'/reveal',json={'kind':'answer'})
    client.post('/api/review/items/'+item['id']+'/attempt',json={'result':'wrong'})
    assert client.get('/api/questions/'+q['id']).json()['state']=='需要巩固'


def test_delete_scrubs_snapshot_and_persistence(client):
    q,body,s=confirmed(client);sess=session(client,q,'delete')
    db.init_db()
    assert client.get('/api/history').json()[0]['id']==q['id']
    assert client.delete('/api/questions/'+q['id']).status_code==200
    assert client.get('/api/history').json()==[]
    assert body['answer'] not in str(client.get('/api/review/sessions/'+sess['id']).json())
    assert db.rows('SELECT * FROM attempts')==[]
    assert client.delete('/api/sources/'+s['id']).status_code==200
    assert not (db.DATA_DIR/'sources'/s['id']).exists()


def test_ai_output_boundary():
    raw={'answer':'答案','correct_parts':'','issues':'错因','evidence':'虚构','check_action':'动作','knowledge_ids':['unknown'],'causes':['语法'],'uncertainty':''}
    r=ai.normalize_analysis(raw,{'work':'原文','missing':False})
    assert r['knowledge_ids']==[] and '证据' in r['uncertainty']
    r=ai.normalize_analysis(raw,{'work':'','missing':True})
    assert not r['issues'] and not r['evidence'] and not r['causes']


def test_invalid_file_and_knowledge_dedup(client):
    assert client.post('/api/sources',data={'name':'bad','learning_date':'2026-09-20'},files={'files':('bad.exe',b'123')}).status_code==400
    q,body,s=confirmed(client)
    client.post('/api/sources/'+s['id']+'/save')
    k=next(k for k in client.get('/api/knowledge').json() if k['id']==body['knowledge_ids'][0])
    assert k['count']==1 and k['wrong_count']==1
    assert len(client.get('/api/history',params={'knowledge':k['id'],'query':'enjoys'}).json())==1

def wait_task(c,tid):
    for _ in range(150):
        t=c.get('/api/tasks/'+tid).json()
        if t['status'] not in ('queued','running'):return t
        time.sleep(.05)
    pytest.fail('task timeout')


def test_recognition_retry_and_analysis_pipeline(client,monkeypatch):
    config.save_settings({'api_key':'sk-fake-key-testing'})
    s=upload(client)
    calls=[]
    def fail(page):
        calls.append(page['id']);raise ai.AIError('模拟额度不足')
    monkeypatch.setattr(ai,'recognize',fail)
    t=client.post('/api/sources/'+s['id']+'/recognize').json()
    assert wait_task(client,t['task_id'])['status']=='failed'
    def success(page):
        calls.append(page['id']);return ([{'label':'1','stem':'计算1+1','work':'1+1=3','missing':False,'warning':''}], '')
    monkeypatch.setattr(ai,'recognize',success)
    t=client.post('/api/sources/'+s['id']+'/recognize').json()
    assert wait_task(client,t['task_id'])['status']=='success'
    q=client.get('/api/sources/'+s['id']).json()['questions'][0]
    t=client.post('/api/sources/'+s['id']+'/recognize').json();wait_task(client,t['task_id'])
    assert len(calls)==2 and len(client.get('/api/sources/'+s['id']).json()['questions'])==1
    def analysis(q,page):
        return {'answer':'2','correct_parts':'列式正确','issues':'加法错误','evidence':'1+1=3','check_action':'复算','causes':['语法'],'knowledge_ids':[domain.NODES[0]['id']],'uncertainty':'','model':'mock-model','prompt_version':'test','knowledge_version':'english-core-1','proposal':{'answer':'2'}}
    monkeypatch.setattr(ai,'analyze',analysis)
    t=client.post('/api/sources/'+s['id']+'/analyze',json={'question_ids':[q['id']]}).json()
    assert wait_task(client,t['task_id'])['status']=='success'
    a=client.get('/api/questions/'+q['id']).json()['analysis']
    assert a['model']=='mock-model' and a['status']=='pending' and a['proposal']=={'answer':'2'}


def test_model_http_contract_and_errors(client,monkeypatch):
    import httpx
    config.save_settings({'api_key':'sk-fake-key-testing'})
    original=httpx.Client
    seen=[]
    def handler(request):
        seen.append(request)
        return httpx.Response(200,json={'choices':[{'message':{'content':'```json\n{"answer":"2"}\n```'}}],'usage':{'total_tokens':12}})
    monkeypatch.setattr(ai.httpx,'Client',lambda **kw:original(transport=httpx.MockTransport(handler),**kw))
    value,usage=ai.chat([{'role':'user','content':'test'}],'qwen-test',thinking=True)
    assert value['answer']=='2' and usage['total_tokens']==12
    payload=db.decode(seen[0].content.decode())
    assert payload['enable_thinking'] is True and payload['stream'] is False
    assert seen[0].url.path.endswith('/chat/completions')
    for code in (401,429,500):
        monkeypatch.setattr(ai.httpx,'Client',lambda **kw:original(transport=httpx.MockTransport(lambda r:httpx.Response(code,json={'error':'sk-fake-key-testing'})),**kw))
        with pytest.raises(ai.AIError) as exc:ai.chat([],'qwen-test')
        assert 'sk-fake' not in str(exc.value)


def test_notes_page_success_and_subject_guard(client,monkeypatch):
    config.save_settings({'api_key':'sk-fake-key-testing'})
    s=upload(client)
    calls=[]
    def read(page):
        calls.append(page['id'])
        return {'text':'Unit 2 Success 学习目标与词组','subject':'物理','content_kind':'notes','questions':[],'warning':'','recognition_version':'page-read-2'}
    monkeypatch.setattr(ai,'recognize',read)
    t=client.post('/api/sources/'+s['id']+'/recognize').json()
    assert wait_task(client,t['task_id'])['status']=='success'
    result=client.get('/api/sources/'+s['id']).json()
    assert result['pages'][0]['subject']=='物理' and result['pages'][0]['text'].startswith('Unit 2')
    assert result['questions']==[]
    t=client.post('/api/sources/'+s['id']+'/recognize').json();wait_task(client,t['task_id'])
    assert len(calls)==1
    q=client.post('/api/sources/'+s['id']+'/questions',json={'page_id':result['pages'][0]['id'],'label':'1','stem':'Translate success','work':'成功'}).json()
    r=client.post('/api/sources/'+s['id']+'/analyze',json={'question_ids':[q['id']]})
    assert r.status_code==400 and '物理' in r.json()['detail']


def test_page_reader_classifies_and_preserves_text(monkeypatch):
    monkeypatch.setattr(ai,'image_content',lambda p:{'type':'text','text':'fixture'})
    monkeypatch.setattr(ai,'chat',lambda *a,**k:({'text':'学习目标','subject':'物理','content_kind':'notes','questions':[],'warning':''},{}))
    result=ai.recognize({'path':'ignored'})
    assert result['content_kind']=='notes' and result['text']=='学习目标'
    monkeypatch.setattr(ai,'chat',lambda *a,**k:({'text':'','questions':[]},{}))
    with pytest.raises(ai.AIError):ai.recognize({'path':'ignored'})


def test_legacy_success_page_refresh_does_not_duplicate(client,monkeypatch):
    config.save_settings({'api_key':'sk-fake-key-testing'})
    q,body,s=confirmed(client)
    page=client.get('/api/sources/'+s['id']).json()['pages'][0]
    db.execute("UPDATE pages SET status='success',recognition_version='' WHERE id=?",(page['id'],))
    monkeypatch.setattr(ai,'recognize',lambda p:{'text':'英语讲义','subject':'物理','content_kind':'notes','questions':[],'warning':'','recognition_version':'page-read-2'})
    t=client.post('/api/sources/'+s['id']+'/recognize').json()
    assert wait_task(client,t['task_id'])['status']=='success'
    result=client.get('/api/sources/'+s['id']).json()
    assert len(result['questions'])==1 and result['pages'][0]['subject']=='物理'


def test_english_analysis_prompt_and_page_confirmation(client,monkeypatch):
    captured=[]
    def chat(messages,*args,**kwargs):
        captured.append(messages[0]['content'][-1]['text'])
        return ({'answer':'reading','correct_parts':'词根正确','issues':'需要动名词','evidence':'read','check_action':'检查enjoy搭配','knowledge_ids':['en_nonfinite'],'causes':['语法'],'uncertainty':''},{})
    monkeypatch.setattr(ai,'chat',chat)
    value=ai.analyze({'stem':'She enjoys ___ (read) books.','work':'read','missing':False})
    assert value['knowledge_version']=='english-core-1' and value['knowledge_ids']==['en_nonfinite']
    assert '高中英语' in captured[0] and '完整原文' in captured[0] and '评分细则' in captured[0]
    source=upload(client);page=source['pages'][0]
    assert client.put('/api/pages/'+page['id']+'/subject',json={'subject':'英语'}).status_code==200
    assert client.get('/api/sources/'+source['id']).json()['pages'][0]['subject']=='英语'


def test_subject_isolation_and_nonenglish_recording(client):
    en,body,ens=confirmed(client)
    b=io.BytesIO();Image.new('RGB',(200,300),'white').save(b,'JPEG')
    r=client.post('/api/sources',data={'name':'数学测试','learning_date':'2026-09-20','subject':'math'},files={'files':('math.jpg',b.getvalue())})
    assert r.status_code==200
    sid=r.json()['id'];wait_task(client,r.json()['task_id'])
    q=client.post('/api/sources/'+sid+'/questions',json={'label':'1','stem':'1+1=?','work':'2'}).json()
    client.put('/api/questions/'+q['id'],json={'version':1,'result':'correct'})
    client.post('/api/sources/'+sid+'/save')
    math_headers={'X-Cuoti-Subject':'math'}
    mh=client.get('/api/history',headers=math_headers).json()
    assert len(mh)==1 and mh[0]['id']==q['id'] and mh[0]['state']=='记录为做对'
    assert client.get('/api/history').json()[0]['id']==en['id']
    assert len(client.get('/api/sources',headers=math_headers).json())==1
    assert client.get('/api/dashboard',headers=math_headers).json()['record_count']==1
    assert client.get('/api/knowledge',headers=math_headers).json()==[]
    assert client.get('/api/review/candidates',headers=math_headers).json()==[]
    assert client.post('/api/sources/'+sid+'/analyze',json={'question_ids':[q['id']]}).status_code==400
    assert client.post('/api/questions/'+q['id']+'/manual-analysis').status_code==400
    assert client.get('/api/sources',headers={'X-Cuoti-Subject':'invalid'}).status_code==400
    db.init_db()
    assert client.get('/api/history',headers=math_headers).json()[0]['id']==q['id']


def test_old_knowledge_analysis_cannot_be_reconfirmed(client):
    q,body,source=confirmed(client)
    db.execute("UPDATE analyses SET knowledge_version='math-core-1',status='stale' WHERE question_id=?",(q['id'],))
    r=client.put('/api/questions/'+q['id']+'/analysis',json=body)
    assert r.status_code==400 and '旧知识库' in r.json()['detail']
    assert client.get('/api/review/candidates').json()==[]


def test_removed_subjects_and_bootstrap(client):
    data=client.get('/api/bootstrap').json()
    assert len(data['subjects'])==7
    assert not {'chemistry','biology'} & {s['id'] for s in data['subjects']}
    assert client.get('/api/history',headers={'X-Cuoti-Subject':'biology'}).status_code==400
